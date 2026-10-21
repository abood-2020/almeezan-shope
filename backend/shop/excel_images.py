"""Read embedded worksheet pictures and map them to product rows."""

from io import BytesIO

from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from shop.remote_images import sniff_image

MAX_WORKBOOK_BYTES = 12 * 1024 * 1024


def _excel_row(anchor):
    marker = getattr(anchor, "_from", None)
    row = getattr(marker, "row", None)
    if row is None:
        return None
    return int(row) + 1


def _workbook(upload):
    payload = upload.read() if hasattr(upload, "read") else upload
    if hasattr(upload, "seek"):
        upload.seek(0)
    if not payload or len(payload) > MAX_WORKBOOK_BYTES or payload[:2] != b"PK":
        raise ValidationError({"file": ["Use an .xlsx workbook of 12 MB or less."]})
    try:
        return load_workbook(BytesIO(payload), data_only=True)
    except Exception as exc:
        raise ValidationError({"file": ["The workbook could not be read."]}) from exc


def _grouped_images(worksheet):
    grouped = {}
    for image in getattr(worksheet, "_images", []):
        excel_row = _excel_row(image.anchor)
        if excel_row is None or excel_row < 2:
            continue
        try:
            data = image._data()
        except Exception:
            grouped.setdefault(excel_row - 2, []).append(b"")
            continue
        grouped.setdefault(excel_row - 2, []).append(data)
    return grouped


def embedded_images_by_row(upload):
    """Return {zero-based sheet row below the header: [image bytes]}."""
    workbook = _workbook(upload)
    try:
        return _grouped_images(workbook.worksheets[0])
    finally:
        workbook.close()


def read_import_workbook(upload):
    """Return non-empty data rows and embedded images aligned to those rows."""
    workbook = _workbook(upload)
    try:
        worksheet = workbook.worksheets[0]
        matrix = [list(values or ()) for values in worksheet.iter_rows(values_only=True)]
        pictures = _grouped_images(worksheet)
    finally:
        workbook.close()
    if not matrix:
        raise ValidationError({"file": ["The workbook is empty."]})
    headers = ["" if cell is None else str(cell).strip() for cell in matrix[0]]
    rows = []
    images = {}
    for offset, values in enumerate(matrix[1:]):
        if not values or all(cell in (None, "") for cell in values):
            continue
        record = {}
        for index, header in enumerate(headers):
            if header and index < len(values):
                record[header] = values[index]
        if pictures.get(offset):
            images[len(rows)] = pictures[offset]
        rows.append(record)
    return rows, images


def preview_embedded_images(upload):
    grouped = embedded_images_by_row(upload)
    images = []
    for index in sorted(grouped):
        pictures = grouped[index]
        excel_row = index + 2
        if len(pictures) > 1:
            images.append({"row": excel_row, "status": "multiple", "count": len(pictures)})
            continue
        kind = sniff_image(pictures[0])
        if not kind:
            images.append({"row": excel_row, "status": "unsupported", "count": 1})
            continue
        import base64
        encoded = base64.b64encode(pictures[0]).decode("ascii") if len(pictures[0]) <= 1_500_000 else ""
        images.append({
            "row": excel_row,
            "status": "ok",
            "count": 1,
            "format": kind,
            "preview": f"data:image/{'jpeg' if kind == 'jpg' else kind};base64,{encoded}" if encoded else "",
        })
    return {"images": images}
