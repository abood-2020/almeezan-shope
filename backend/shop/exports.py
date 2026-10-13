"""Excel workbooks for saved orders and invoice records.

Order values come from the submitted snapshot. Current catalog prices, price lists,
and discounts are not read back into a historical export.
"""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

PINK = "B41F71"
INK = "2D2730"
HEADER_FILL = PatternFill("solid", fgColor=PINK)
TOTAL_FILL = PatternFill("solid", fgColor="F8E8F1")
WHITE = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
BODY = Font(name="Calibri", color=INK, size=11)
BOLD = Font(name="Calibri", bold=True, color=INK, size=11)
THIN = Border(
    left=Side(style="thin", color="E9DFE5"),
    right=Side(style="thin", color="E9DFE5"),
    top=Side(style="thin", color="E9DFE5"),
    bottom=Side(style="thin", color="E9DFE5"),
)
STATUS = {
    "review": ("قيد المراجعة", "Under review"),
    "confirmed": ("مؤكد", "Confirmed"),
    "complete": ("مكتمل", "Completed"),
    "cancelled": ("ملغى", "Cancelled"),
}


def normalize_language(value):
    token = str(value or "").strip().lower().replace("_", "-")
    if token.startswith("en"):
        return "en"
    return "ar"


def export_language(request):
    """Use the requested UI language. Arabic is the fallback when it is missing or unknown."""
    explicit = request.query_params.get("lang")
    if explicit is not None and str(explicit).strip():
        return normalize_language(explicit)
    header = request.META.get("HTTP_ACCEPT_LANGUAGE", "")
    for part in header.split(","):
        token = part.split(";", 1)[0].strip().lower()
        if token.startswith("en"):
            return "en"
        if token.startswith("ar"):
            return "ar"
    return "ar"


def xlsx_response(payload, filename):
    from django.http import HttpResponse

    safe = "".join(char for char in filename if char.isalnum() or char in "._-") or "export.xlsx"
    response = HttpResponse(
        payload,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{safe}"'
    response["Cache-Control"] = "no-store"
    return response


def sheet_text(value):
    text = "" if value is None else str(value)
    if text.lstrip()[:1] in ("=", "+", "-", "@"):
        return "'" + text
    return text


def split_option(label):
    """Read color and size from the saved option label. Do not look up the live product."""
    raw = label or ""
    names, _separator, size = raw.partition(" · ")
    color_ar, _slash, color_en = names.partition(" / ")
    if not color_en:
        color_en = color_ar
    return color_ar.strip(), color_en.strip(), size.strip()


def _line_total(line):
    if line.line_total is not None:
        return line.line_total
    return line.unit_price * line.quantity


def _prepare(ws: Worksheet, english: bool):
    ws.sheet_view.rightToLeft = not english
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4


def _header(ws, titles, english):
    alignment = Alignment(horizontal="left" if english else "right", vertical="center", wrap_text=True)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(titles))}2"
    ws.freeze_panes = "A3"
    ws.row_dimensions[2].height = 22
    for index, title in enumerate(titles, start=1):
        cell = ws.cell(2, index, sheet_text(title))
        cell.font = WHITE
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return alignment


def _widths(ws, widths):
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_title(ws, text, columns):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = ws.cell(1, 1, sheet_text(text))
    cell.font = Font(name="Calibri", bold=True, color=PINK, size=16)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28


def _money(cell, value):
    if value is None or value == "":
        cell.value = None
        return
    cell.value = float(Decimal(value))
    cell.number_format = '#,##0.00'
    cell.alignment = Alignment(horizontal="right", vertical="center")


def order_workbook(order, language, setting):
    english = normalize_language(language) == "en"
    brand = setting.name_en if english else setting.name_ar
    wb = Workbook()
    summary = wb.active
    summary.title = "Order Summary" if english else "ملخص الطلب"
    items = wb.create_sheet("Order Items" if english else "بنود الطلب")
    _summary_sheet(summary, order, english, brand)
    _items_sheet(items, order, english, brand)
    _embed_logo(summary, setting)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _summary_sheet(ws, order, english, brand):
    _prepare(ws, english)
    titles = [en if english else ar for ar, en in (("الحقل", "Field"), ("القيمة", "Value"))]
    _write_title(ws, f"{brand} — {order.number}", 2)
    alignment = _header(ws, titles, english)
    status = STATUS.get(order.status, (order.status, order.status))
    price_list = order.price_list_name_en if english else order.price_list_name_ar
    trader_name = order.trader_name_en if english and order.trader_name_en else order.trader_name_ar
    rows = [
        ("Order Number", "رقم الطلب", order.number),
        ("Order Date", "تاريخ الطلب", order.ordered_on),
        ("Order Status", "حالة الطلب", status[1] if english else status[0]),
        ("Trader Name", "اسم التاجر", trader_name),
        ("Business Name", "الاسم التجاري", trader_name),
        ("Email", "البريد الإلكتروني", order.trader_email),
        ("Currency", "العملة", order.currency_id),
        ("Price List", "قائمة الأسعار", price_list),
        ("Notes", "الملاحظات", order.note),
        ("Total Quantity", "إجمالي الكمية", sum(line.quantity for line in order.lines.all())),
        ("Order Total", "إجمالي الطلب", order.total),
    ]
    for offset, (en_label, ar_label, value) in enumerate(rows, start=3):
        label = ws.cell(offset, 1, sheet_text(en_label if english else ar_label))
        label.font = BOLD
        label.alignment = alignment
        label.border = THIN
        cell = ws.cell(offset, 2)
        cell.font = BODY
        cell.border = THIN
        cell.alignment = alignment
        if en_label == "Order Total":
            _money(cell, value)
            label.fill = TOTAL_FILL
            cell.fill = TOTAL_FILL
            label.font = BOLD
            cell.font = BOLD
        elif en_label == "Total Quantity":
            cell.value = int(value)
            cell.number_format = "0"
        elif en_label == "Order Date":
            cell.value = value
            cell.number_format = "YYYY-MM-DD"
        else:
            cell.value = sheet_text(value)
    _widths(ws, (28, 42))
    ws.auto_filter.ref = f"A2:B{2 + len(rows)}"
    ws.oddFooter.center.text = sheet_text(brand)


def _items_sheet(ws, order, english, brand):
    _prepare(ws, english)
    columns = [
        ("الرقم", "No."),
        ("كود المنتج", "Product Code"),
        ("الباركود", "Barcode"),
        ("اسم المنتج", "Product Name"),
        ("التصنيف", "Category"),
        ("اللون", "Color"),
        ("المقاس", "Size"),
        ("وحدة البيع", "Unit"),
        ("الكمية", "Quantity"),
        ("السعر الأساسي", "Base Price"),
        ("نسبة الخصم", "Discount %"),
        ("سعر الوحدة", "Unit Price"),
        ("إجمالي السطر", "Line Total"),
        ("العملة الأصلية", "Original Currency"),
        ("السعر الأصلي", "Original Price"),
    ]
    titles = [en if english else ar for ar, en in columns]
    _write_title(ws, f"{brand} — {order.number}", len(titles))
    alignment = _header(ws, titles, english)
    lines = list(order.lines.all())
    for index, line in enumerate(lines, start=1):
        row = index + 2
        color_ar, color_en, size = split_option(line.option_label)
        category = line.category_path_en if english else line.category_path_ar
        unit = line.unit_en if english else line.unit_ar
        name = line.product_name_en if english and line.product_name_en else line.product_name_ar
        values = [
            index,
            line.product_code,
            "",
            name,
            category,
            color_en if english else color_ar,
            size,
            unit,
            line.quantity,
            line.base_price,
            line.discount_percent,
            line.unit_price,
            _line_total(line),
            line.original_currency,
            line.original_price,
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row, column)
            cell.font = BODY
            cell.border = THIN
            cell.alignment = alignment
            if column in {1, 9}:
                cell.value = int(value or 0)
                cell.number_format = "0"
            elif column in {10, 12, 13, 15}:
                _money(cell, value)
            elif column == 11:
                if value is None:
                    cell.value = None
                else:
                    cell.value = float(Decimal(value))
                    cell.number_format = "0.00"
            elif column in {2, 3}:
                cell.value = None if value in (None, "") else sheet_text(value)
                cell.number_format = "@"
            else:
                cell.value = sheet_text(value)
    total_row = len(lines) + 3
    label = ws.cell(total_row, 12, sheet_text("Order Total" if english else "إجمالي الطلب"))
    amount = ws.cell(total_row, 13)
    label.font = BOLD
    label.fill = TOTAL_FILL
    label.border = THIN
    amount.font = BOLD
    amount.fill = TOTAL_FILL
    amount.border = THIN
    _money(amount, order.total)
    _widths(ws, (8, 18, 16, 36, 28, 18, 12, 14, 12, 18, 14, 16, 16, 16, 16))
    ws.auto_filter.ref = f"A2:{get_column_letter(len(titles))}{max(total_row - 1, 2)}"
    ws.oddFooter.center.text = sheet_text(brand)


def invoice_workbook(invoice, language, setting):
    """Invoice records store a number, date, amount, and file. They have no item lines, status, or currency snapshot."""
    english = normalize_language(language) == "en"
    brand = setting.name_en if english else setting.name_ar
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice" if english else "الفاتورة"
    _prepare(ws, english)
    titles = ["Field" if english else "الحقل", "Value" if english else "القيمة"]
    _write_title(ws, f"{brand} — {invoice.number}", 2)
    alignment = _header(ws, titles, english)
    trader = invoice.trader
    rows = [
        ("Invoice Number", "رقم الفاتورة", invoice.number),
        ("Invoice Date", "تاريخ الفاتورة", invoice.issued_on),
        ("Trader", "التاجر", (trader.name_en if english else trader.name_ar) if trader else ""),
        ("Recorded amount", "القيمة المسجلة", invoice.amount),
        ("Currency", "العملة", ""),
        ("File name", "اسم الملف", invoice.file_name),
    ]
    for offset, (en_label, ar_label, value) in enumerate(rows, start=3):
        label = ws.cell(offset, 1, sheet_text(en_label if english else ar_label))
        label.font = BOLD
        label.alignment = alignment
        label.border = THIN
        cell = ws.cell(offset, 2)
        cell.font = BODY
        cell.border = THIN
        cell.alignment = alignment
        if en_label == "Recorded amount":
            _money(cell, value)
            label.fill = TOTAL_FILL
            cell.fill = TOTAL_FILL
        elif en_label == "Invoice Date":
            cell.value = value
            cell.number_format = "YYYY-MM-DD"
        else:
            cell.value = sheet_text(value)
    _widths(ws, (28, 36))
    ws.auto_filter.ref = f"A2:B{2 + len(rows)}"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _embed_logo(ws, setting):
    if not setting.logo:
        return
    name = setting.logo.name.lower()
    if not name.endswith((".png", ".jpg", ".jpeg")):
        return
    try:
        from openpyxl.drawing.image import Image as XLImage

        image = XLImage(setting.logo.path)
        image.height = 42
        image.width = 42
        ws.add_image(image, "A1")
    except Exception:
        return
