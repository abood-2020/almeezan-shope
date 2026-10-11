from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from shop.models import Category, Order, OrderLine, Product
from shop.services import create_order


def workbook(response):
    return load_workbook(BytesIO(response.content))


class ExcelExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_initial_data")

    def setUp(self):
        self.client = APIClient()

    def login(self, username, password="demo123"):
        self.client.post("/api/auth/login/", {"username": username, "password": password}, format="json")

    def staff(self):
        self.client.post("/api/admin/auth/login/", {"username": "admin", "password": "demo123"}, format="json")

    def test_historical_order_excel_keeps_saved_prices_and_variants(self):
        line = OrderLine.objects.get(order__number="RW-2026-1042", product_id=1)
        line.product_code = "00123"
        line.save(update_fields=["product_code"])
        Order.objects.filter(number="RW-2026-1042").update(note="=1+1")
        Product.objects.filter(pk=1).update(base_price_ils="1.00")
        self.login("trader")
        response = self.client.get("/api/orders/RW-2026-1042/excel/?lang=en")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Order_RW-2026-1042.xlsx", response["Content-Disposition"])
        book = workbook(response)
        self.assertEqual(book.sheetnames, ["Order Summary", "Order Items"])
        self.assertFalse(book["Order Summary"].sheet_view.rightToLeft)
        summary = {book["Order Summary"].cell(row, 1).value: book["Order Summary"].cell(row, 2).value for row in range(3, 16)}
        self.assertEqual(summary["Order Number"], "RW-2026-1042")
        self.assertEqual(summary["Order Total"], 1914)
        self.assertEqual(summary["Total Quantity"], 18)
        self.assertIsNone(summary["Price List"])
        self.assertIsNone(summary["Email"])
        self.assertEqual(summary["Notes"], "'=1+1")
        items = book["Order Items"]
        self.assertEqual(items["A2"].value, "No.")
        self.assertEqual(items["B2"].value, "Product Code")
        self.assertEqual(items["M2"].value, "Line Total")
        self.assertEqual(items["B3"].value, "00123")
        self.assertEqual(items["B3"].number_format, "@")
        self.assertEqual(items["C3"].value, None)
        self.assertEqual(items["F3"].value, "Ivory")
        self.assertEqual(items["G3"].value, "M")
        self.assertEqual(items["L3"].value, 82)
        self.assertEqual(items["L4"].value, 155)
        self.assertEqual(items["I3"].value, 12)
        self.assertEqual(items["I4"].value, 6)
        self.assertNotEqual(items["F3"].value, items["F4"].value)

    def test_new_order_snapshots_are_stable_and_arabic_export_is_rtl(self):
        trader = User.objects.get(username="trader").trader
        order = create_order(trader, "", [
            {"product_id": 1, "option": "عاجي / Ivory · M", "qty": 12},
            {"product_id": 1, "option": "رملي / Sand · M", "qty": 12},
        ], True)
        self.assertEqual(order.price_list_name_ar, "تجار فلسطين")
        self.assertEqual(order.lines.first().category_path_ar, "رجالي")
        Category.objects.filter(name_ar="رجالي").update(name_ar="CHANGED")
        Product.objects.filter(pk=1).update(base_price_ils="9.00")
        self.login("trader")
        response = self.client.get(f"/api/orders/{order.number}/excel/?lang=ar")
        self.assertEqual(response.status_code, 200)
        book = workbook(response)
        self.assertEqual(book.sheetnames, ["ملخص الطلب", "بنود الطلب"])
        self.assertTrue(book["بنود الطلب"].sheet_view.rightToLeft)
        summary = {book["ملخص الطلب"].cell(row, 1).value: book["ملخص الطلب"].cell(row, 2).value for row in range(3, 16)}
        self.assertEqual(summary["قائمة الأسعار"], "تجار فلسطين")
        self.assertEqual(summary["البريد الإلكتروني"], "trader@example.com")
        items = book["بنود الطلب"]
        self.assertEqual(items["A2"].value, "الرقم")
        self.assertEqual(items["B2"].value, "كود المنتج")
        self.assertEqual(items["E3"].value, "رجالي")
        self.assertEqual(items["E4"].value, "رجالي")
        self.assertEqual(items["F3"].value, "عاجي")
        self.assertEqual(items["F4"].value, "رملي")
        self.assertEqual(items["I3"].value, 12)
        self.assertEqual(items["I4"].value, 12)
        self.assertNotEqual(float(items["L3"].value), 9)

    def test_language_follows_the_request_and_arabic_is_the_fallback(self):
        self.login("trader")
        english = workbook(self.client.get("/api/orders/RW-2026-1042/excel/?lang=EN-us"))
        arabic = workbook(self.client.get("/api/orders/RW-2026-1042/excel/?lang=fr"))
        header = workbook(self.client.get("/api/orders/RW-2026-1042/excel/", HTTP_ACCEPT_LANGUAGE="en-US,en;q=0.8"))
        fallback = workbook(self.client.get("/api/orders/RW-2026-1042/excel/"))

        def total(book, sheet, label):
            worksheet = book[sheet]
            for row in range(3, 20):
                if worksheet.cell(row, 1).value == label:
                    return worksheet.cell(row, 2).value
            return None

        self.assertEqual(english.sheetnames, ["Order Summary", "Order Items"])
        self.assertFalse(english["Order Items"].sheet_view.rightToLeft)
        self.assertEqual(arabic.sheetnames, ["ملخص الطلب", "بنود الطلب"])
        self.assertTrue(arabic["ملخص الطلب"].sheet_view.rightToLeft)
        self.assertEqual(header.sheetnames[0], "Order Summary")
        self.assertEqual(fallback.sheetnames[0], "ملخص الطلب")
        self.assertEqual(total(english, "Order Summary", "Order Total"), total(arabic, "ملخص الطلب", "إجمالي الطلب"))
        self.assertEqual(english["Order Items"]["L3"].value, arabic["بنود الطلب"]["L3"].value)

    def test_export_permissions(self):
        guest = self.client.get("/api/orders/RW-2026-1042/excel/")
        self.assertEqual(guest.status_code, 403)
        self.login("export")
        other = self.client.get("/api/orders/RW-2026-1042/excel/")
        self.assertEqual(other.status_code, 404)
        self.staff()
        staff = self.client.get("/api/admin/orders/RW-2026-1042/excel/?lang=en")
        self.assertEqual(staff.status_code, 200)
        staff_book = workbook(staff)
        summary = {staff_book["Order Summary"].cell(row, 1).value: staff_book["Order Summary"].cell(row, 2).value for row in range(3, 16)}
        self.assertEqual(summary["Order Total"], 1914)

    def test_invoice_excel_uses_saved_amount_without_invented_lines(self):
        self.login("trader")
        response = self.client.get("/api/invoices/INV-2026-087/excel/?lang=en")
        self.assertEqual(response.status_code, 200)
        book = workbook(response)
        self.assertEqual(book.sheetnames, ["Invoice"])
        values = {book["Invoice"].cell(row, 1).value: book["Invoice"].cell(row, 2).value for row in range(3, 10)}
        self.assertEqual(values["Invoice Number"], "INV-2026-087")
        self.assertEqual(values["Recorded amount"], 456)
        self.assertIsNone(values["Currency"])
        self.assertNotIn("Order Number", values)
        arabic = workbook(self.client.get("/api/invoices/INV-2026-087/excel/?lang=ar"))
        self.assertEqual(arabic.sheetnames, ["الفاتورة"])
        self.assertTrue(arabic["الفاتورة"].sheet_view.rightToLeft)
        arabic_values = {arabic["الفاتورة"].cell(row, 1).value: arabic["الفاتورة"].cell(row, 2).value for row in range(3, 10)}
        self.assertEqual(arabic_values["رقم الفاتورة"], "INV-2026-087")
        self.assertEqual(arabic_values["القيمة المسجلة"], 456)
        self.login("export")
        self.assertEqual(self.client.get("/api/invoices/INV-2026-087/excel/").status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get("/api/invoices/INV-2026-087/excel/").status_code, 403)


@override_settings(MEDIA_ROOT="media-test")
class StoreContentTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_initial_data")

    def setUp(self):
        self.client = APIClient()

    def test_public_defaults_and_staff_brand_update(self):
        public = self.client.get("/api/settings/")
        self.assertEqual(public.status_code, 200)
        self.assertEqual(public["Cache-Control"], "no-store")
        self.assertEqual(public.data["brand"]["name_ar"], "رِواق")
        self.assertEqual(public.data["brand"]["name_en"], "Riwaq")
        self.assertEqual(len(public.data["faqs"]), 4)
        self.assertTrue(public.data["sections"][0]["visible"])
        self.client.post("/api/admin/auth/login/", {"username": "admin", "password": "demo123"}, format="json")
        saved = self.client.patch("/api/admin/settings/", {"name_ar": "بيت الأزياء", "name_en": "Fashion House"}, format="json")
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.data["brand"]["name_ar"], "بيت الأزياء")
        self.client.logout()
        again = self.client.get("/api/settings/")
        self.assertEqual(again.data["brand"]["name_en"], "Fashion House")

    def test_traders_and_guests_cannot_edit_content(self):
        self.assertEqual(self.client.patch("/api/admin/settings/", {"name_en": "Nope"}, format="json").status_code, 401)
        self.client.post("/api/auth/login/", {"username": "trader", "password": "demo123"}, format="json")
        self.assertEqual(self.client.patch("/api/admin/settings/", {"name_en": "Nope"}, format="json").status_code, 403)
        self.assertEqual(self.client.post("/api/admin/faqs/", {"question_ar": "س", "question_en": "Q", "answer_ar": "ج", "answer_en": "A"}, format="json").status_code, 403)

    def test_section_visibility_faq_and_validation(self):
        self.client.post("/api/admin/auth/login/", {"username": "admin", "password": "demo123"}, format="json")
        hidden = self.client.patch("/api/admin/content/sections/", {"sections": [{"key": "benefits", "visible": False}]}, format="json")
        self.assertEqual(hidden.status_code, 200)
        self.assertFalse(next(row for row in hidden.data["sections"] if row["key"] == "benefits")["visible"])
        created = self.client.post("/api/admin/faqs/", {
            "question_ar": "سؤال جديد", "question_en": "New question",
            "answer_ar": "جواب", "answer_en": "Answer",
        }, format="json")
        self.assertEqual(created.status_code, 201)
        new_id = created.data["faqs"][-1]["id"]
        ordered = self.client.put("/api/admin/faqs/order/", {"ids": [new_id]}, format="json")
        self.assertEqual(ordered.status_code, 200)
        self.assertEqual(ordered.data["faqs"][0]["id"], new_id)
        inactive = self.client.patch(f"/api/admin/faqs/{new_id}/", {"is_active": False}, format="json")
        self.assertEqual(inactive.status_code, 200)
        self.client.logout()
        public = self.client.get("/api/settings/")
        self.assertFalse(any(item["id"] == new_id for item in public.data["faqs"]))
        self.assertFalse(next(row for row in public.data["sections"] if row["key"] == "benefits")["visible"])
        self.client.post("/api/admin/auth/login/", {"username": "admin", "password": "demo123"}, format="json")
        markup = self.client.patch("/api/admin/settings/", {"name_en": "<script>alert(1)</script>"}, format="json")
        self.assertEqual(markup.status_code, 400)
        upload = self.client.patch("/api/admin/settings/", {"logo": SimpleUploadedFile("logo.txt", b"hello", content_type="text/plain")}, format="multipart")
        self.assertEqual(upload.status_code, 400)
        video = self.client.patch("/api/admin/settings/", {"hero_video_url": "http://example.com/clip.mp4"}, format="json")
        self.assertEqual(video.status_code, 400)
        contact = self.client.patch("/api/admin/settings/", {
            "company_whatsapp_number": "970599111222",
            "email": "hello@riwaq.test",
            "instagram_url": "https://instagram.com/riwaq",
        }, format="json")
        self.assertEqual(contact.status_code, 200)
        self.assertEqual(contact.data["contact"]["whatsapp"], "970599111222")
        self.assertEqual(contact.data["contact"]["email"], "hello@riwaq.test")

    def test_partial_saves_keep_hidden_store_fields(self):
        self.client.post("/api/admin/auth/login/", {"username": "admin", "password": "demo123"}, format="json")
        before = self.client.get("/api/admin/settings/").data
        saved = self.client.patch("/api/admin/settings/", {"name_en": "Riwaq Atelier"}, format="json")
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.data["brand"]["name_en"], "Riwaq Atelier")
        self.assertEqual(saved.data["seo"]["title_en"], before["seo"]["title_en"])
        self.assertEqual(saved.data["contact"]["whatsapp"], before["contact"]["whatsapp"])
        self.assertEqual(saved.data["hero"]["video_url"], before["hero"]["video_url"])
        hero = self.client.patch("/api/admin/content/sections/", {"sections": [{"key": "hero", "title_en": "Hello traders"}]}, format="json")
        self.assertEqual(hero.status_code, 200)
        self.assertEqual(next(row for row in hero.data["sections"] if row["key"] == "hero")["title_en"], "Hello traders")
        benefits = next(row for row in hero.data["sections"] if row["key"] == "benefits")
        before_benefits = next(row for row in before["sections"] if row["key"] == "benefits")
        self.assertEqual(benefits["visible"], before_benefits["visible"])
        self.assertEqual(benefits["title_ar"], before_benefits["title_ar"])
